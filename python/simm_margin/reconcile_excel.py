"""Independent Python recomputation of the Excel workbook, LibreOffice recalculation and comparison. SYNTHETIC data.

Run: py -3 -m simm_margin.reconcile_excel [--quick]   (needs LibreOffice for the recalculation step).
The workbook builder (excel/build_workbook.py) types the Python reference values onto Python_Ref using reference() here;
this module then compares the recalculated cells of the workbook with the same reference.
"""
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from . import attribution as attr
from . import backtest as bt
from . import columns as C
from . import config
from . import params as params_mod
from . import parameter_io as pio
from . import schedule_im as sched
from . import simm as simm_mod

# ===== CONFIG (user inputs) =====
EXCEL_DIR = config.ROOT / "excel"
BUILDER = EXCEL_DIR / "build_workbook.py"
XLSX = EXCEL_DIR / "simm_margin_workbook.xlsx"
REPORT = config.OUT_DIR / "excel_reconciliation.json"
OPTION_INPUTS = config.EXCEL_INPUT_DIR / "option_checks.json"
SOFFICE_CANDIDATES = [r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
AUTHOR = "Deepak Chaudhary"
CHAMPION_MODEL, SIMM_MODEL = "hs_im_eu_1plus3", "simm_style"
SPLIT_1D, SPLIT_10D = "1d", "10d_nonoverlap"
SERIES_COL_MODEL, SERIES_COL_SPLIT = C.MODEL, "split"
QUICK_DAYS_1D, QUICK_DAYS_10D = 700, 60
LAST_WINDOW = 250
IR_CCYS = ("USD", "EUR", "GBP", "JPY", "MXN")
FX_CCYS = ("EUR", "GBP", "JPY", "MXN")
FX_PAIRS = ("EURUSD", "GBPUSD", "USDJPY", "USDMXN")
SWAPTION_CHECK_TRADE, FXOPT_CHECK_TRADE = "SWPT_USD_1Yx5Y", "FXO_GBPUSD_STRADDLE_C"
TOL_REL, TOL_USD, TOL_LR, TOL_P, TOL_COUNT, TOL_SHAPLEY = 1e-9, 0.01, 1e-9, 1e-10, 0.0, 0.01
TOL_ANALYTIC = 1e-3                  # analytic closed form vs central-difference bump in the CRIF
KIND_REL, KIND_ABS, KIND_TEXT = "rel", "abs", "text"
CASES = (
    dict(name="base", param_set=config.SIMM_VERSION, scale={}, p_exc=bt.P_EXC),
    dict(name="usd10y_x1.5_params2506_p2pct", param_set=config.SIMM_VERSION_PRIOR, scale={"IRS_USD_10Y": 1.5}, p_exc=0.02),
    dict(name="concentration_stress_cr_above_1", param_set=config.SIMM_VERSION,
         scale={"IRS_USD_10Y": 20000.0, "SWPT_USD_1Yx5Y": 10000.0, "FXF_EURUSD_6M": 10000.0, "FXO_GBPUSD_STRADDLE_C": 8000.0},
         p_exc=bt.P_EXC),
)
LO_TIMEOUT_S = 600
# ===== END CONFIG =====


def key(*parts) -> str:
    """Join key parts with a pipe; the builder uses the same function for the Python_Ref keys."""
    return "|".join(str(p) for p in parts)


# ---------------------------------------------------------------- inputs
def option_check_inputs() -> dict:
    """Closed-form check inputs (forward swap rate, annuity, vol, discount factors) at the valuation date."""
    if OPTION_INPUTS.exists():
        return json.loads(OPTION_INPUTS.read_text(encoding="utf-8"))
    return compute_option_inputs()


def compute_option_inputs() -> dict:
    from . import instruments, market_history, pricing
    hist = market_history.generate_history(quick=True)
    row = hist.market.row(len(hist) - 1)
    tr = instruments.sample_portfolio().set_index(C.TRADE_ID)
    t1 = tr.loc[SWAPTION_CHECK_TRADE]
    fwd, ann = row.curve(t1[C.CCY]).forward_swap(t1[C.EXPIRY], t1[C.MATURITY])
    sig = pricing.ir_vol(row, t1[C.CCY], t1[C.EXPIRY], t1[C.MATURITY] - t1[C.EXPIRY])[0]
    t2 = tr.loc[FXOPT_CHECK_TRADE]
    b, q, tt = t2[C.CCY], t2[C.CCY2], t2[C.EXPIRY]
    return {
        "swaption": dict(trade_id=SWAPTION_CHECK_TRADE, notional=float(t1[C.NOTIONAL]), direction=float(t1[C.DIRECTION]),
                         forward=float(fwd[0]), annuity=float(ann[0]), strike=float(t1[C.STRIKE]), sigma_bp=float(sig * 1e4),
                         expiry=float(t1[C.EXPIRY]), fx_to_usd=float(row.fx_spot[t1[C.CCY]][0]), payer=1.0),
        "fxoption": dict(trade_id=FXOPT_CHECK_TRADE, notional=float(t2[C.NOTIONAL]), direction=float(t2[C.DIRECTION]),
                         forward=float(pricing.fx_forward_rate(row, b, q, tt)[0]), strike=float(t2[C.STRIKE]),
                         sigma=float(pricing.fx_vol(row, b + q, tt)[0]), expiry=float(tt),
                         df_quote=float(row.curve(q).df(tt)[0, 0]), df_base=float(row.curve(b).df(tt)[0, 0]),
                         spot_base=float(row.fx_spot[b][0]), fx_quote_to_usd=float(row.fx_spot[q][0])),
    }


def write_option_inputs() -> dict:
    d = compute_option_inputs()
    OPTION_INPUTS.parent.mkdir(parents=True, exist_ok=True)
    OPTION_INPUTS.write_text(json.dumps(d, indent=1), encoding="utf-8")
    return d


def _series(df: pd.DataFrame, model: str, split: str, n_last) -> pd.DataFrame:
    s = df[(df[SERIES_COL_MODEL] == model) & (df[SERIES_COL_SPLIT] == split)][[C.DATE, C.LOSS, C.VAR_FORECAST]]
    s = s.sort_values(C.DATE).reset_index(drop=True)
    return s.iloc[-n_last:].reset_index(drop=True) if n_last else s


def load_inputs(quick: bool = False) -> dict:
    """Everything the workbook types from Python, as DataFrames and dicts (unmodified by any case)."""
    out = config.OUT_DIR
    series = pd.read_csv(out / "backtest_series.csv")
    rows = pio.read_csv(pio.SIMM_PARAMS_FILE)
    sets = (config.SIMM_VERSION, config.SIMM_VERSION_PRIOR)
    schedule = [r for r in pio.read_csv(pio.SCHEDULE_FILE) if sched.SOURCE_US in r["source_doc"]]
    consts = [r for r in pio.read_csv(pio.REG_CONSTANTS_FILE)
              if r["constant"] in (sched.CONST_GROSS, sched.CONST_NGR) and r["jurisdiction"] == "US"]
    tl = pd.read_csv(config.PARAM_DIR / pio.TRAFFIC_LIGHT_FILE)
    return dict(
        crif=pd.read_csv(out / "crif_last.csv"), portfolio=pd.read_csv(out / "portfolio_pv.csv"),
        subsets=pd.read_csv(out / "attribution_subset_ims.csv"),
        scenario=pd.read_csv(out / "attribution_scenario_day.csv"),
        s1d=_series(series, CHAMPION_MODEL, SPLIT_1D, QUICK_DAYS_1D if quick else 0),
        s10d=_series(series, SIMM_MODEL, SPLIT_10D, QUICK_DAYS_10D if quick else 0),
        param_rows=[r for r in rows if r[C.PARAM_SET] in sets], param_sets=sets,
        schedule_rows=schedule, constants=consts, bcbs=tl, fx_spot=sched.default_fx_spot(),
        options=option_check_inputs(), quick=quick, data_source=config.DATA_SOURCE)


def apply_case(inp: dict, case: dict) -> dict:
    """Copy of the inputs with the case applied: trade scaling rewrites CRIF, notional, PV and option-check notionals."""
    out = dict(inp)
    crif, port = inp["crif"].copy(), inp["portfolio"].copy()
    opts = json.loads(json.dumps(inp["options"]))
    for tid, f in case["scale"].items():
        m = crif[C.CRIF_TRADE_ID] == tid
        crif.loc[m, [C.AMOUNT, C.AMOUNT_USD]] = crif.loc[m, [C.AMOUNT, C.AMOUNT_USD]] * f
        pm = port[C.TRADE_ID] == tid
        port.loc[pm, [C.NOTIONAL, "pv"]] = port.loc[pm, [C.NOTIONAL, "pv"]] * f
        for o in opts.values():
            if o["trade_id"] == tid:
                o["notional"] *= f
    out.update(crif=crif, portfolio=port, options=opts, case=case, param_set=case["param_set"], p_exc=case["p_exc"])
    return out


# ---------------------------------------------------------------- Python reference
def _bd_index(bd: pd.DataFrame) -> dict:
    return {(r.level, r.risk_class, r.margin_type, r.bucket, r.risk_factor): r.value for r in bd.itertuples(index=False)}


def reference(inp: dict) -> dict:
    """Python values (recomputed through the project modules) for every quantity the workbook replicates."""
    p = params_mod.load(inp["param_set"])
    crif = inp["crif"]
    res = simm_mod.simm(crif, p)
    ws = simm_mod.weighted_sensitivities(crif, p)
    ix = _bd_index(res.breakdown)
    L = simm_mod
    ref = {}
    tenors = config.IR_TENORS
    ws_ir = {(r.bucket, r.tenor): r.WS for r in ws[ws[C.RISK_CLASS] == C.RC_IR].itertuples(index=False)}
    ws_fx = {r.risk_factor: (r.WS, r.CR) for r in ws[ws[C.RISK_CLASS] == C.RC_FX].itertuples(index=False)}
    for ccy in IR_CCYS:
        for t in tenors:
            ref[key("ws", "IR", ccy, t)] = float(ws_ir.get((ccy, t), 0.0))
            ref[key("vr", "IR", ccy, t)] = float(ix.get((L.LEVEL_FACTOR, C.RC_IR, C.MT_VEGA, ccy, t), 0.0))
            ref[key("c", "IR", ccy, t)] = float(ix.get((L.LEVEL_FACTOR, C.RC_IR, C.MT_CURVATURE, ccy, t), 0.0))
        for mt in (C.MT_DELTA, C.MT_VEGA, C.MT_CURVATURE):
            ref[key("k", "IR", mt, ccy)] = float(ix.get((L.LEVEL_K, C.RC_IR, mt, ccy, ""), 0.0))
            ref[key("s", "IR", mt, ccy)] = float(ix.get((L.LEVEL_S, C.RC_IR, mt, ccy, ""), 0.0))
        for mt in (C.MT_DELTA, C.MT_VEGA):
            ref[key("cr", "IR", mt, ccy)] = float(ix.get((L.LEVEL_CR, C.RC_IR, mt, ccy, ""), 1.0))
    for ccy in FX_CCYS:
        ref[key("ws", "FX", ccy)] = float(ws_fx.get(ccy, (0.0, 1.0))[0])
        ref[key("cr", "FX", C.MT_DELTA, ccy)] = float(ws_fx.get(ccy, (0.0, 1.0))[1])
    for pair in FX_PAIRS:
        ref[key("sigma", pair)] = float(p.fx_sigma(pair[:3], pair[3:]))
        ref[key("vr", "FX", pair)] = float(ix.get((L.LEVEL_FACTOR, C.RC_FX, C.MT_VEGA, L.FX_BUCKET, pair), 0.0))
        ref[key("c", "FX", pair)] = float(ix.get((L.LEVEL_FACTOR, C.RC_FX, C.MT_CURVATURE, L.FX_BUCKET, pair), 0.0))
    for mt in (C.MT_DELTA, C.MT_VEGA, C.MT_CURVATURE):
        ref[key("k", "FX", mt)] = float(ix.get((L.LEVEL_K, C.RC_FX, mt, L.FX_BUCKET, ""), 0.0))
    for rc in (C.RC_IR, C.RC_FX):
        ref[key("theta", rc)] = float(ix.get(("theta", rc, C.MT_CURVATURE, "", ""), 0.0))
        ref[key("lambda", rc)] = float(ix.get(("lambda", rc, C.MT_CURVATURE, "", ""), 0.0))
        ref[key("im", rc)] = float(res.by_risk_class[rc])
        for mt in (C.MT_DELTA, C.MT_VEGA, C.MT_CURVATURE):
            ref[key("margin", rc, mt)] = float(res.by_margin_type[(rc, mt)])
    ref[key("total")] = float(res.total)
    for i, t in enumerate(tenors):
        ref[key("sf", t)] = float(p.sf(p.tenor_days())[i])
    # schedule IM
    port = inp["portfolio"]
    sc = sched.schedule_im(port, port["pv"].to_numpy(), fx_spot=inp["fx_spot"])
    ref[key("sched", "gross")], ref[key("sched", "ngr")], ref[key("sched", "net")] = sc["gross"], sc["ngr"], sc["net"]
    for r in sc["by_row"].itertuples(index=False):
        ref[key("sched", "rate", r.trade_id)] = float(r.rate)
        ref[key("sched", "gross_im", r.trade_id)] = float(r.gross_im)
    # backtests
    pe = inp["p_exc"]
    for tag, df in (("1d", inp["s1d"]), ("10d", inp["s10d"])):
        e = bt.exceptions(df[C.LOSS].to_numpy(), df[C.VAR_FORECAST].to_numpy())
        ku, ch = bt.kupiec_pof(int(e.sum()), len(e), pe), bt.christoffersen(e, pe)
        ref[key("bt", tag, "n")], ref[key("bt", tag, "x")] = float(len(e)), float(e.sum())
        ref[key("bt", tag, "LR_uc")], ref[key("bt", tag, "p_uc")] = ku["LR"], ku["p_value"]
        for k_ in ("n00", "n01", "n10", "n11", "LR_ind", "p_ind", "LR_cc", "p_cc"):
            ref[key("bt", tag, k_)] = float(ch[k_])
        ref[key("bt", tag, "zone")] = bt.basel_zone(int(e.sum()), len(e), pe)
        if tag == "1d":
            x250 = int(e[-LAST_WINDOW:].sum())
            ref[key("bt", tag, "x250")] = float(x250)
            ref[key("bt", tag, "zone250")] = bt.basel_zone(x250, LAST_WINDOW, pe)
    # attribution
    sub = inp["subsets"]
    drivers = list(attr.DRIVERS)
    vals = {frozenset(d for d in drivers if r[d] == 1): r["im"] for _, r in sub.iterrows()}
    shap = attr.shapley_values(vals)
    im0, im1 = vals[frozenset()], vals[frozenset(drivers)]
    steps = [frozenset(drivers[:k]) for k in range(len(drivers) + 1)]
    for i, d in enumerate(drivers):
        ref[key("attr", "shapley", d)] = float(shap[d])
        ref[key("attr", "oaat", d)] = float(vals[frozenset([d])] - im0)
        ref[key("attr", "bridge", d)] = float(vals[steps[i + 1]] - vals[steps[i]])
    ref[key("attr", "substeps_sum")] = ref[key("attr", "bridge", "market_move")]
    ref[key("attr", "dim")] = float(im1 - im0)
    ref[key("attr", "shapley_sum")] = float(sum(shap.values()))
    ref[key("attr", "oaat_residual")] = float(im1 - im0 - sum(vals[frozenset([d])] - im0 for d in drivers))
    # closed-form option checks (the reference is the CRIF amount produced by the bump-and-reprice sensitivities)
    o = inp["options"]
    cr = crif
    def amt(tid, rt):
        return float(cr[(cr[C.CRIF_TRADE_ID] == tid) & (cr[C.RISK_TYPE] == rt)][C.AMOUNT_USD].sum())
    pvs = inp["portfolio"].set_index(C.TRADE_ID)["pv"]
    ref[key("opt", "bachelier_price")] = float(pvs[o["swaption"]["trade_id"]])
    ref[key("opt", "gk_price")] = float(pvs[o["fxoption"]["trade_id"]])
    ref[key("opt", "bachelier_vega_amount")] = amt(o["swaption"]["trade_id"], C.RISK_IRVOL)
    ref[key("opt", "gk_vega_amount")] = amt(o["fxoption"]["trade_id"], C.RISK_FXVOL)
    ref[key("opt", "gk_delta_amount")] = amt(o["fxoption"]["trade_id"], C.RISK_FX)
    return ref


def tolerance(k: str):
    """(kind, tolerance) for a reference key."""
    f = k.split("|")
    if f[0] == "bt":
        if f[2] in ("zone", "zone250"):
            return KIND_TEXT, 0.0
        if f[2] in ("n", "x", "x250", "n00", "n01", "n10", "n11"):
            return KIND_ABS, TOL_COUNT
        if f[2].startswith("p_"):
            return KIND_ABS, TOL_P
        return KIND_ABS, TOL_LR
    if f[0] == "attr":
        return KIND_ABS, TOL_SHAPLEY
    if f[0] == "opt":
        return (KIND_REL, TOL_REL) if f[1].endswith("price") else (KIND_REL, TOL_ANALYTIC)
    if f[0] in ("total", "im") or (f[0] == "sched" and f[1] in ("gross", "net")):
        return KIND_ABS, TOL_USD
    if f[0] == "margin":
        return KIND_REL, TOL_REL
    return KIND_REL, TOL_REL


def block_of(k: str) -> str:
    f = k.split("|")
    if f[0] in ("ws", "sigma", "sf"):
        return "weighted_sensitivities" if f[0] == "ws" else "parameters_derived"
    if f[0] in ("k", "s", "cr", "vr", "c", "theta", "lambda"):
        return "bucket_K_S_CR_VR"
    if f[0] in ("margin", "im"):
        return "margins"
    if f[0] == "total":
        return "total_im"
    return {"sched": "schedule_im", "bt": "backtest", "attr": "attribution", "opt": "analytic_checks"}[f[0]]


# ---------------------------------------------------------------- LibreOffice
def find_soffice():
    for c in SOFFICE_CANDIDATES:
        if Path(c).exists():
            return c
    return None


def load_builder():
    spec = importlib.util.spec_from_file_location("simm_build_workbook", BUILDER)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def recalc(src: Path, outdir: Path, soffice: str) -> Path:
    shutil.rmtree(outdir, ignore_errors=True)
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_")).as_uri()
    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(outdir), str(src)], check=True, capture_output=True, timeout=LO_TIMEOUT_S)
    return outdir / src.name


def patch_author(path: Path, author: str = AUTHOR):
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                x = data.decode("utf-8")
                for tag in ("dc:creator", "cp:lastModifiedBy"):
                    if re.search(rf"<{tag}>.*?</{tag}>", x, flags=re.S):
                        x = re.sub(rf"<{tag}>.*?</{tag}>", f"<{tag}>{author}</{tag}>", x, flags=re.S)
                    elif re.search(rf"<{tag}\s*/>", x):
                        x = re.sub(rf"<{tag}\s*/>", f"<{tag}>{author}</{tag}>", x)
                    else:
                        x = x.replace("</cp:coreProperties>", f"<{tag}>{author}</{tag}></cp:coreProperties>")
                data = x.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


# ---------------------------------------------------------------- comparison
def _diff(kind, a, b) -> float:
    if kind == KIND_TEXT:
        return 0.0 if str(a) == str(b) else 1.0
    a, b = float(a), float(b)
    return abs(a - b) / max(1.0, abs(b)) if kind == KIND_REL else abs(a - b)


def compare(label: str, wbv, meta: dict, ref: dict) -> dict:
    issues, worst, counts = [], {}, {}
    for chk in meta["checks"]:
        k = chk["key"]
        v = wbv[chk["sheet"]][chk["cell"]].value
        kind, tol = tolerance(k)
        if v is None or (isinstance(v, str) and v.startswith("#")) or (kind != KIND_TEXT and isinstance(v, str)):
            issues.append(f"{k}: bad cell value {v!r}")
            continue
        d = _diff(kind, v, ref[k])
        blk = block_of(k)
        counts[blk] = counts.get(blk, 0) + 1
        worst[blk] = max(worst.get(blk, 0.0), d)
        if d > tol:
            issues.append(f"{k}: excel {v!r} python {ref[k]!r} diff {d:.3g} > {tol:g}")
    ck = wbv["Checks"]
    all_cell = ck[meta["allpass_cell"]].value
    if all_cell != 1:
        issues.append(f"workbook Checks all-pass cell is {all_cell!r}")
    n_in = sum(1 for r in range(meta["checks_first"], meta["checks_last"] + 1) if ck[f"{meta['pass_col']}{r}"].value == 1)
    n_all = meta["checks_last"] - meta["checks_first"] + 1
    if n_in != n_all:
        issues.append(f"in-workbook checks passed {n_in} of {n_all}")
    bad = [f"{ws.title}!{c.coordinate}" for ws in wbv for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
    errs = len(bad)
    if errs:
        issues.append(f"{errs} error cells in the recalculated workbook, first: {bad[:8]}")
    return dict(case=label, passed=not issues, issues=issues, n_compared=len(meta["checks"]), compared_by_block=counts,
                worst_diffs=worst, workbook_checks=f"{n_in}/{n_all}", error_cells=errs)


# ---------------------------------------------------------------- driver
def build_and_reconcile(quick: bool = False, xlsx_out=None, report_out=None, cases=None) -> dict:
    t0 = time.time()
    soffice = find_soffice()
    if soffice is None:
        msg = "LibreOffice not found; install it (winget install TheDocumentFoundation.LibreOffice). Skipping recalculation."
        print(msg)
        return dict(all_passed=None, skipped=True, reason=msg, data_source=config.DATA_SOURCE)
    xlsx_out = Path(xlsx_out or XLSX)
    report_out = Path(report_out or REPORT)
    builder = load_builder()
    base = load_inputs(quick)
    work = Path(tempfile.mkdtemp(prefix="simm_xl_"))
    results, shipped, meta0 = [], None, None
    for i, case in enumerate(cases or CASES):
        inp = apply_case(base, case)
        ref = reference(inp)
        src = work / f"case{i}.xlsx"
        meta = builder.build(src, inp, ref)
        out = recalc(src, work / f"out{i}", soffice)
        wbv = openpyxl.load_workbook(out, data_only=True)
        res = compare(case["name"], wbv, meta, ref)
        results.append(res)
        print(("PASS" if res["passed"] else "FAIL"), case["name"], res["n_compared"], "checks", res["issues"][:3])
        if i == 0:
            shipped, meta0 = out, meta
    xlsx_out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(shipped, xlsx_out)
    patch_author(xlsx_out)
    meta0 = {k: v for k, v in meta0.items()}
    xlsx_out.with_suffix(".meta.json").write_text(json.dumps(meta0, indent=1, default=str), encoding="utf-8")
    worst = {}
    for r in results:
        for k, v in r["worst_diffs"].items():
            worst[k] = max(worst.get(k, 0.0), v)
    report = dict(data_source=config.DATA_SOURCE, all_passed=all(r["passed"] for r in results), quick=quick,
                  n_cases=len(results), worst_diffs=worst, cases=results, n_formulas=meta0.get("n_formulas"),
                  sheets=meta0.get("sheets"), n_checks_per_case=results[0]["n_compared"],
                  libreoffice=str(soffice), runtime_s=round(time.time() - t0, 1))
    report_out.parent.mkdir(parents=True, exist_ok=True)
    report_out.write_text(json.dumps(report, indent=1, default=float), encoding="utf-8")
    shutil.rmtree(work, ignore_errors=True)
    print("all passed:", report["all_passed"], "runtime", report["runtime_s"], "s")
    return report


def main():
    if not OPTION_INPUTS.exists():
        write_option_inputs()
    build_and_reconcile(quick="--quick" in sys.argv)


if __name__ == "__main__":
    main()
