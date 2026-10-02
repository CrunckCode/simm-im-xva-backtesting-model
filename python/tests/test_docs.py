"""Guards on the generated documentation (docx, markdown, README); reads files only and never rewrites python/outputs."""
import re
import sys
import zipfile
from pathlib import Path

import pytest

# ===== CONFIG (user inputs) =====
ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
STEM = "SIMM_IM_xVA_Backtesting_Documentation"
DOCX, MD, README = DOCS / f"{STEM}.docx", DOCS / f"{STEM}.md", ROOT / "README.md"
AUTHOR = "Deepak Chaudhary"
EM_DASH, EN_DASH = chr(0x2014), chr(0x2013)
DASH_SUBSTITUTE = re.compile(r"(?<=\w)--(?=\w)|\s--\s|" + EM_DASH + "|" + EN_DASH)
# a citation is a parenthetical that starts with one of the existing sources
CITATION = re.compile(r"\((?:ISDA SIMM|ISDA Governance|ISDA Risk Data|BCBS-IOSCO|BCBS 22|12 CFR 45\.8|12 CFR Part 45|EU 2016/2251|"
                      r"Delegated Regulation|Kupiec 1995|Christoffersen 1998|SR 11-7)[^()]*\)")
DISCLAIMER_TOKENS = ("synthetic", "not licensed by isda", "no compliance claim")
# ===== END CONFIG =====

pytestmark = pytest.mark.skipif(not DOCX.exists(), reason="docs not built")
sys.path.insert(0, str(DOCS))


def _docx_text():
    from docx import Document
    d = Document(str(DOCX))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_files_exist():
    assert DOCX.exists() and MD.exists() and README.exists()


def test_no_em_en_or_double_hyphen_dashes():
    for name, text in (("docx", _docx_text()), ("md", MD.read_text(encoding="utf-8")), ("readme", README.read_text(encoding="utf-8"))):
        assert not DASH_SUBSTITUTE.search(text), f"dash substitute in {name}: {DASH_SUBSTITUTE.search(text)}"
        assert chr(0xFFFD) not in text


def test_heading_styles_have_no_color():
    xml = zipfile.ZipFile(DOCX).read("word/styles.xml").decode("utf-8")
    blocks = re.findall(r"<w:style [^>]*w:styleId=\"Heading\d\".*?</w:style>", xml, flags=re.S)
    assert blocks
    for block in blocks:
        assert "<w:color" not in block and "themeColor" not in block


def test_author_properties():
    from docx import Document
    cp = Document(str(DOCX)).core_properties
    assert cp.author == AUTHOR and cp.last_modified_by == AUTHOR


def test_counts_match_markdown():
    import build_docs as bd
    assert bd.md_stats(MD.read_text(encoding="utf-8")) == bd.docx_stats(DOCX)


def test_image_paths_exist():
    text = MD.read_text(encoding="utf-8")
    paths = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    assert len(paths) >= 12
    for p in paths:
        assert (DOCS / p).resolve().exists(), p


def test_at_most_one_citation_per_paragraph():
    for para in re.split(r"\n\s*\n", MD.read_text(encoding="utf-8")):
        assert len(CITATION.findall(para)) <= 1, para[:150]


def test_no_unresolved_placeholders():
    for path in (MD, README):
        text = path.read_text(encoding="utf-8")
        assert "@@" not in text
        assert not re.findall(r"\bnan\b", text)


def test_disclaimer_comes_first():
    for path in (MD, README):
        body = re.sub(r"^---\n.*?\n---\n", "", path.read_text(encoding="utf-8"), count=1, flags=re.S)
        head = body[:1500].lower()
        for token in DISCLAIMER_TOKENS:
            assert token in head, (path.name, token)


def test_verification_status_is_reported():
    text = MD.read_text(encoding="utf-8")
    assert "Original text NOT obtained" in text and "12 CFR 237.8" in text and "CFTC 23.154" in text
    assert "ISDA credit, equity and commodity" in text


def test_required_sections_present():
    heads = re.findall(r"^# (.+)$", MD.read_text(encoding="utf-8"), flags=re.M)
    assert len(heads) == 19


def test_excel_section_never_claims_missing_numbers():
    text = MD.read_text(encoding="utf-8")
    recon = ROOT / "python" / "outputs" / "excel_reconciliation.json"
    if not recon.exists():
        assert "being finalized" in text
