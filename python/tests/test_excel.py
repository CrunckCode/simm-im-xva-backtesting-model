"""Workbook tests: plain formatting, author, formula/typed separation, static conclusions, LibreOffice reconciliation. SYNTHETIC data."""
import json
import re
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pytest

from simm_margin import reconcile_excel as rx

# ===== CONFIG (user inputs) =====
CALC_SHEETS = ("SIMM_Delta", "SIMM_Vega", "SIMM_Curvature", "SIMM_Total", "Schedule_IM", "Checks")
STATIC_SHEETS = ("README", "Conclusions")
MIN_CASES = 3
BAD_DASH = re.compile("[–—]|--")
BLACK = ("FF000000", "00000000")
# ===== END CONFIG =====

needs_lo = pytest.mark.skipif(rx.find_soffice() is None, reason="LibreOffice not installed")


@pytest.fixture(scope="module")
def tmp_dir():
    with tempfile.TemporaryDirectory(prefix="simm_xl_test_") as d:
        yield Path(d)


@pytest.fixture(scope="module")
def report(tmp_dir):
    return rx.build_and_reconcile(quick=True, xlsx_out=tmp_dir / "wb.xlsx", report_out=tmp_dir / "report.json")


@pytest.fixture(scope="module")
def wb(report, tmp_dir):
    return openpyxl.load_workbook(tmp_dir / "wb.xlsx")


@pytest.fixture(scope="module")
def meta(report, tmp_dir):
    return json.loads((tmp_dir / "wb.meta.json").read_text())


@needs_lo
def test_reconciliation_passes(report, tmp_dir):
    assert report["all_passed"], [c["issues"][:5] for c in report["cases"] if not c["passed"]]
    assert report["n_cases"] >= MIN_CASES
    assert json.loads((tmp_dir / "report.json").read_text())["data_source"]
    assert all(c["error_cells"] == 0 for c in report["cases"])


@needs_lo
def test_plain_formatting(wb):
    for ws in wb:
        assert ws.sheet_properties.tabColor is None, ws.title
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                assert c.fill is None or c.fill.fill_type in (None, "none"), (ws.title, c.coordinate)
                color = c.font.color
                assert color is None or color.type != "rgb" or color.rgb in BLACK, (ws.title, c.coordinate)


@needs_lo
def test_author_metadata(wb, tmp_dir):
    assert wb.properties.creator == rx.AUTHOR
    assert wb.properties.lastModifiedBy == rx.AUTHOR
    with zipfile.ZipFile(tmp_dir / "wb.xlsx") as z:
        assert rx.AUTHOR in z.read("docProps/core.xml").decode("utf-8")


@needs_lo
def test_calculation_sheets_are_formulas_only(wb):
    for name in CALC_SHEETS:
        for row in wb[name].iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                assert not (isinstance(v, (int, float)) and not isinstance(v, bool)), (name, c.coordinate, v)
                if isinstance(v, str) and not v.startswith("="):
                    assert c.font.bold, (name, c.coordinate, v)       # typed text on calc sheets is a bold label


@needs_lo
def test_mixed_sheets_type_numbers_only_in_declared_ranges(wb, meta):
    for name, ranges in meta["typed"].items():
        if name not in ("Sensitivities", "Backtest", "Attribution"):
            continue
        ok = set()
        for rng in ranges:
            for row in wb[name][rng]:
                for c in (row if isinstance(row, tuple) else (row,)):
                    ok.add(c.coordinate)
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    assert c.coordinate in ok, (name, c.coordinate, c.value)


@needs_lo
def test_conclusions_and_readme_static(wb):
    for name in STATIC_SHEETS:
        n = 0
        for row in wb[name].iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), (name, c.coordinate)
                n += c.value is not None
        assert n > 0


@needs_lo
def test_no_dashes_in_workbook(wb):
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    assert not BAD_DASH.search(c.value), (ws.title, c.coordinate)


@needs_lo
def test_checks_sheet_all_pass_cell(wb, meta):
    assert meta["n_formulas"] > 5000
    assert wb["Checks"][meta["allpass_cell"]].value is not None


@needs_lo
def test_report_has_blocks(report):
    for blk in ("weighted_sensitivities", "margins", "total_im", "schedule_im", "backtest", "attribution"):
        assert blk in report["worst_diffs"], blk


def test_python_reference_is_self_consistent():
    inp = rx.apply_case(rx.load_inputs(quick=True), rx.CASES[0])
    ref = rx.reference(inp)
    assert ref[rx.key("total")] > 0
    parts = ref[rx.key("im", "IR")] + ref[rx.key("im", "FX")]
    assert ref[rx.key("total")] <= parts + 1e-6
    assert abs(ref[rx.key("attr", "shapley_sum")] - ref[rx.key("attr", "dim")]) < 1e-6
